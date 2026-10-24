"""PantheonAI Word document collector.

This module scans ``.docx`` files, extracts their text content, and wraps the
result in a structured :class:`DocumentObject` ready for Markdown generation
and AI indexing.
"""

from __future__ import annotations

from hashlib import sha256
from pathlib import Path
from typing import Iterable

try:
    from docx import Document as DocxDocument
except Exception as exc:  # pragma: no cover - dependency guarded at runtime
    DocxDocument = None  # type: ignore[assignment]
    DOCX_IMPORT_ERROR = exc
else:  # pragma: no cover - import sentinel for happy path
    DOCX_IMPORT_ERROR = None

from pantheon_ai.models.document_object import DocumentObject


class WordCollector:
    """Collect ``.docx`` files from a folder tree."""

    def __init__(self, input_folder: str | Path):
        self.input_folder = Path(input_folder)
        if not self.input_folder.exists():
            raise FileNotFoundError(f"Input folder not found: {self.input_folder}")
        if not self.input_folder.is_dir():
            raise NotADirectoryError(f"Input folder is not a directory: {self.input_folder}")

    def scan(self) -> list[Path]:
        """Return all supported Word files under the input folder."""

        return sorted(
            path
            for path in self.input_folder.rglob("*.docx")
            if path.is_file()
        )

    def read_file(self, file_path: Path) -> DocumentObject:
        """Read and extract a single Word file."""

        raw = file_path.read_bytes()
        document = self._open_document(file_path)
        text = self._extract_text(document)
        relative = file_path.relative_to(self.input_folder)

        return DocumentObject(
            source_type="Word",
            title=file_path.stem,
            name=file_path.stem,
            filename=file_path.name,
            relative_path=str(relative.parent),
            absolute_path=str(file_path),
            text=text,
            encoding="binary/docx",
            size=len(raw),
            page_count=0,
            sha256=sha256(raw).hexdigest(),
            metadata={
                "paragraphs": len(document.paragraphs),
                "tables": len(document.tables),
                "extensions": [".docx"],
            },
        )

    def collect(self) -> list[DocumentObject]:
        """Collect all supported Word documents."""

        documents: list[DocumentObject] = []
        for file_path in self.scan():
            try:
                documents.append(self.read_file(file_path))
            except Exception as exc:  # pragma: no cover - collector must continue
                print(f"[ERROR] {file_path}: {exc}")
        return documents

    def collect_iter(self) -> Iterable[DocumentObject]:
        """Yield Word documents one by one."""

        for file_path in self.scan():
            try:
                yield self.read_file(file_path)
            except Exception as exc:  # pragma: no cover - collector must continue
                print(f"[ERROR] {file_path}: {exc}")

    def _open_document(self, file_path: Path):
        if DocxDocument is None:
            raise ImportError(
                "python-docx is required for Word document support. Install it with 'pip install python-docx'."
            ) from DOCX_IMPORT_ERROR
        return DocxDocument(str(file_path))

    def _extract_text(self, document) -> str:
        parts: list[str] = []

        for paragraph in document.paragraphs:
            text = self._normalize(paragraph.text)
            if text:
                parts.append(text)

        for table in document.tables:
            table_rows: list[str] = []
            for row in table.rows:
                cells = [self._normalize(cell.text) for cell in row.cells]
                if any(cells):
                    table_rows.append(" | ".join(cells))
            if table_rows:
                parts.append("\n".join(table_rows))

        return self._normalize("\n\n".join(parts))

    @staticmethod
    def _normalize(value: str) -> str:
        return " ".join(value.replace("\r\n", "\n").replace("\r", "\n").split())


if __name__ == "__main__":
    collector = WordCollector(Path("input") / "docs")
    docs = collector.collect()

    print(f"Found {len(docs)} Word files")
    for doc in docs:
        print(f"- {doc.relative_path}/{doc.filename} ({doc.size} bytes)")
